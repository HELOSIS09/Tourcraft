# 





from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

# Palette
NAVY = RGBColor(25, 31, 43)
WINE = RGBColor(111, 28, 53)
GOLD = RGBColor(202, 158, 75)
CREAM = RGBColor(248, 245, 239)
WHITE = RGBColor(255, 255, 255)
MUTED = RGBColor(103, 108, 118)
GREEN = RGBColor(45, 125, 93)

def bg(slide, color=CREAM):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color

def textbox(slide, x, y, w, h, text="", size=20, color=NAVY, bold=False,
            font="Aptos", align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    r.font.name = font
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color
    return box

def rect(slide, x, y, w, h, fill, radius=False, line=None):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
        Inches(x), Inches(y), Inches(w), Inches(h)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line if line else fill
    return shape

def title(slide, kicker, main, num):
    textbox(slide, 0.65, 0.35, 10.8, 0.35, kicker.upper(), 11, GOLD, True)
    textbox(slide, 0.65, 0.78, 11.7, 0.7, main, 29, NAVY, True)
    textbox(slide, 12.15, 0.45, 0.5, 0.35, f"{num:02d}", 12, WINE, True, align=PP_ALIGN.RIGHT)

def bullet_box(slide, x, y, w, h, heading, bullets, accent=WINE):
    rect(slide, x, y, w, h, WHITE, True)
    rect(slide, x, y, 0.07, h, accent)
    textbox(slide, x+0.25, y+0.18, w-0.45, 0.38, heading, 16, NAVY, True)
    box = slide.shapes.add_textbox(Inches(x+0.25), Inches(y+0.65), Inches(w-0.45), Inches(h-0.8))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    for i, b in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = b
        p.level = 0
        p.font.name = "Aptos"
        p.font.size = Pt(13)
        p.font.color.rgb = NAVY
        p.space_after = Pt(9)
        p.text = "• " + b

# Slide 1
s = prs.slides.add_slide(prs.slide_layouts[6])
bg(s)
title(s, "Business Pitch | TourCraft AI", "A tourism-operations problem I have experienced first-hand", 1)

# Founder card
rect(s, 0.65, 1.75, 3.55, 4.95, NAVY, True)
textbox(s, 0.95, 2.05, 2.9, 0.45, "APURVA SWARNKAAR", 20, WHITE, True)
textbox(s, 0.95, 2.52, 2.9, 0.35, "Founder • Tourism & Hospitality", 12, GOLD, True)
textbox(s, 0.95, 3.05, 2.9, 1.25,
        "B.Voc Hospitality & Tourism Management\nBHU • RGSC\nOperations internship experience in a Vietnam DMC",
        13, WHITE)
rect(s, 0.95, 4.65, 2.85, 0.72, WINE, True)
textbox(s, 1.12, 4.82, 2.5, 0.35, "8.5K+ LinkedIn network", 15, WHITE, True, align=PP_ALIGN.CENTER)
textbox(s, 0.95, 5.65, 2.9, 0.7,
        "Tourism-industry connections + direct exposure to B2B travel operations",
        11.5, WHITE)

bullet_box(s, 4.55, 1.75, 3.85, 4.95, "THE PROBLEM", [
    "Travel agents still spend hours building itineraries, quotations and operational follow-ups manually.",
    "Information is scattered across sheets, messages, supplier documents and CRMs.",
    "Small agencies often lack affordable tools that understand real travel operations."
], WINE)

bullet_box(s, 8.65, 1.75, 4.0, 4.95, "WHY NOW", [
    "India recorded 32.83M Indian national departures and 4.29B domestic tourist visits in 2025.",
    "Tourism is large enough for specialized B2B workflow software.",
    "AI can automate repetitive work without replacing the agent's relationships."
], GOLD)

textbox(s, 8.95, 6.2, 3.4, 0.3, "Source: Ministry of Tourism, India Tourism Data", 8.5, MUTED)

# Slide 2
s = prs.slides.add_slide(prs.slide_layouts[6])
bg(s)
title(s, "The Product", "TourCraft AI turns travel operations into one workflow", 2)

# workflow
steps = [
    ("01", "INPUT", "Client needs\n+ budget\n+ dates"),
    ("02", "PLAN", "AI itinerary\n+ destination\n+ activities"),
    ("03", "QUOTE", "Hotels + services\n+ pricing logic"),
    ("04", "OPERATE", "Confirmations\n+ vouchers\n+ tracking"),
]
xs = [0.7, 3.85, 7.0, 10.15]
for i, (n, h, b) in enumerate(steps):
    rect(s, xs[i], 1.8, 2.55, 2.15, WHITE, True)
    rect(s, xs[i]+0.18, 2.0, 0.52, 0.52, WINE if i != 2 else GOLD, True)
    textbox(s, xs[i]+0.18, 2.09, 0.52, 0.25, n, 10, WHITE, True, align=PP_ALIGN.CENTER)
    textbox(s, xs[i]+0.18, 2.7, 2.1, 0.3, h, 12, WINE, True)
    textbox(s, xs[i]+0.18, 3.08, 2.15, 0.65, b, 12.5, NAVY, True)
    if i < 3:
        textbox(s, xs[i]+2.63, 2.62, 0.45, 0.4, "→", 23, GOLD, True, align=PP_ALIGN.CENTER)

bullet_box(s, 0.7, 4.35, 3.8, 2.15, "CORE MVP", [
    "AI itinerary generator",
    "Budget-aware hotel/service suggestions",
    "Professional quotation + PDF generation"
], WINE)
bullet_box(s, 4.78, 4.35, 3.8, 2.15, "B2B OPERATIONS", [
    "Booking/confirmation tracker",
    "Supplier coordination workflow",
    "On-trip execution status"
], GOLD)
bullet_box(s, 8.86, 4.35, 3.8, 2.15, "LONG-TERM PLATFORM", [
    "Agent CRM + supplier marketplace",
    "Analytics & productivity dashboard",
    "AI operations assistant"
], GREEN)

# Slide 3
s = prs.slides.add_slide(prs.slide_layouts[6])
bg(s)
title(s, "Traction & Business Model", "I am not starting from zero — I already have distribution", 3)

# traction stats
stats = [
    ("200+", "Travel-agent relationships / access"),
    ("8.5K+", "LinkedIn tourism network"),
    ("1", "DMC/supplier relationship covering Vietnam, Cambodia & Laos"),
]
for i, (big, small) in enumerate(stats):
    x = 0.7 + i*4.05
    rect(s, x, 1.65, 3.7, 1.55, NAVY, True)
    textbox(s, x+0.2, 1.9, 3.3, 0.5, big, 27, GOLD, True, align=PP_ALIGN.CENTER)
    textbox(s, x+0.25, 2.48, 3.2, 0.45, small, 10.5, WHITE, False, align=PP_ALIGN.CENTER)

bullet_box(s, 0.7, 3.55, 3.8, 2.75, "WHO PAYS?", [
    "Travel agencies & tour operators",
    "Small DMCs and B2B travel businesses",
    "Later: supplier-side subscriptions / transaction revenue"
], WINE)
bullet_box(s, 4.78, 3.55, 3.8, 2.75, "REVENUE MODEL", [
    "Freemium / trial for lead generation",
    "Pro subscription for operations & AI",
    "B2B plans + optional transaction/service fees"
], GOLD)
bullet_box(s, 8.86, 3.55, 3.8, 2.75, "GO-TO-MARKET", [
    "Pilot with 5–10 travel agents",
    "Use existing tourism network for feedback",
    "Convert successful pilots into paid accounts"
], GREEN)

textbox(s, 0.75, 6.65, 12.0, 0.3,
        "Key advantage: founder-market fit + direct access to early users + operational knowledge.", 11, WINE, True, align=PP_ALIGN.CENTER)

# Slide 4
s = prs.slides.add_slide(prs.slide_layouts[6])
bg(s)
title(s, "The Ask & 24-Hour Execution", "Funding will turn a promising prototype into a testable business", 4)

# left: 24 hour
bullet_box(s, 0.65, 1.65, 5.8, 4.95, "NEXT 24 HOURS — BUILD THE BUSINESS", [
    "Lock the MVP: itinerary → quotation → hotel/service logic → PDF",
    "Add a simple agent login + saved enquiries",
    "Create a live supplier/destination database for 3–5 destinations",
    "Deploy a working demo on a public domain",
    "Onboard 5 pilot agents and collect real feedback",
    "Measure: time saved, quotes created, repeat usage and conversion"
], WINE)

# right resources
bullet_box(s, 6.75, 1.65, 5.9, 2.25, "RESOURCES NEEDED TO SCALE", [
    "Cloud hosting + database + AI/API credits",
    "Product/UI developer support",
    "Tourism supplier partnerships & verified rates",
    "Sales/customer-success support"
], GOLD)

bullet_box(s, 6.75, 4.15, 5.9, 2.45, "WHAT I WANT TO PROVE", [
    "Agents will use it repeatedly — not just like the demo.",
    "It can reduce quotation/itinerary workload significantly.",
    "A small paid pilot can become recurring B2B revenue.",
    "The product can expand destination-by-destination."
], GREEN)

rect(s, 6.95, 6.25, 5.5, 0.42, NAVY, True)
textbox(s, 7.05, 6.33, 5.3, 0.22, "ASK: funding + mentorship + pilot introductions", 11, WHITE, True, align=PP_ALIGN.CENTER)

# Save
path = "/mnt/data/TourCraft_AI_4_Slide_Business_Pitch.pptx"
prs.save(path)
print(path)
